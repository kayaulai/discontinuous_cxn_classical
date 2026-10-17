#!/usr/bin/env python3
"""
add_tags.py  --  add tags to the adverb / final-particle data

    python add_tags.py exp1_site_data.xlsx exp1_selected_df_tagged.csv

reads  the data made by convert_exp1.py (an .xlsx with a "pairs" sheet, or a .csv with the same headings)
       the tagged file (columns dc, sfp and cat_long)
writes exp1_site_data_tagged.csv  -- one row per pair: word1, word2, sym_score, asym_score,
                                     asym_score_2, attr_pair_tag
       (these are the site's standard column names, so the file can be opened directly on the
        upload-your-own site or given to build_site.py without a "columns" setting)

TAGS
  Each pair's cat_long text is cut wherever it has " - " or " / " (a space on each side), and every
  piece is one tag:
      "Modal - Declarative / Information structure - Focus - Cause presentation"
   -> Modal | Declarative | Information structure | Focus | Cause presentation
  Pieces joined without spaces ("Certainty/Expected", "Non-existence", "Question-word") are left whole.
  Extra spaces are tidied, and spellings that differ only in capital letters are made the same
  (the commonest spelling wins).  All tags of a pair go in one column, separated by " | ".

Needs:  pip install openpyxl     (only if the data file is an .xlsx)
"""
import argparse, csv, re, sys
from collections import Counter
from pathlib import Path

# ============================================================ SETTINGS: edit here
# headings in the data file made by convert_exp1.py  ->  the site's standard column names
DATA_HEADINGS = {
    "ADV": "word1", "FP": "word2", "PMI": "sym_score",
    "ADV -> FP": "asym_score", "FP -> ADV": "asym_score_2",
    # a data file that already uses the standard names works too:
    "word1": "word1", "word2": "word2", "sym_score": "sym_score",
    "asym_score": "asym_score", "asym_score_2": "asym_score_2",
}
# columns of the tagged file
TAG_COL_FIRST, TAG_COL_SECOND, TAG_COL_TEXT = "dc", "sfp", "cat_long"
# what counts as a divider between tags (needs a space on each side; the dash can be - or a long dash)
DIVIDER = re.compile(r"\s+(?:-|–|—|/)\s+")
# name of the tag column in the result ("attr_pair_" makes it a pair filter on the site)
OUT_TAG_COLUMN = "attr_pair_tag"
TAG_JOIN = " | "
# ============================================================ end of settings

OUT_COLUMNS = ["word1", "word2", "sym_score", "asym_score", "asym_score_2"]


def die(msg):
    sys.exit("\nERROR: " + msg + "\n")


def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:      # utf-8-sig drops the byte-order mark
        rd = csv.DictReader(f)
        return rd.fieldnames or [], list(rd)


def read_data(path):
    """The pairs, as a list of dicts keyed by the site's standard names (values stay text)."""
    p = Path(path)
    if not p.exists():
        die("Data file not found: %s" % p)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            die("openpyxl is needed to read an .xlsx file:   pip install openpyxl")
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb["pairs"] if "pairs" in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
        if not rows:
            die("The data file is empty.")
        heads, body = [str(h).strip() for h in rows[0]], rows[1:]
        recs = [dict(zip(heads, r)) for r in body if any(str(c).strip() for c in r)]
    else:
        heads, recs = read_csv(p)
        heads = [h.strip() for h in heads]
    mapping = {h: DATA_HEADINGS[h] for h in heads if h in DATA_HEADINGS}
    for need in ("word1", "word2"):
        if need not in mapping.values():
            die("Could not find the %s column in %s.\n  Headings found: %s\n  Add yours to DATA_HEADINGS at the top of this script."
                % (need, p.name, ", ".join(heads)))
    out = []
    for r in recs:
        out.append({std: ("%s" % r[h]).strip() for h, std in mapping.items()})
    return out


def split_tags(text):
    text = re.sub(r"\s+", " ", (text or "")).strip()      # "Quantification  - Non-existence" -> one space
    return [t.strip() for t in DIVIDER.split(text) if t.strip()]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("data", help="exp1_site_data.xlsx (or a .csv with the same headings)")
    ap.add_argument("tagged", help="exp1_selected_df_tagged.csv")
    ap.add_argument("--out", help="result file (default: <data file name>_tagged.csv next to the data file)")
    args = ap.parse_args()

    data = read_data(args.data)
    t_cols, t_rows = read_csv(args.tagged)
    for need in (TAG_COL_FIRST, TAG_COL_SECOND, TAG_COL_TEXT):
        if need not in t_cols:
            die("The tagged file has no column '%s'.\n  Columns found: %s" % (need, ", ".join(t_cols)))

    # ---- tags per pair
    raw_tags, dup = {}, 0
    for r in t_rows:
        key = (r[TAG_COL_FIRST].strip(), r[TAG_COL_SECOND].strip())
        if not key[0] or not key[1]:
            continue
        if key in raw_tags:
            dup += 1
            continue
        raw_tags[key] = split_tags(r[TAG_COL_TEXT])

    # ---- make spellings that differ only in capitals the same (commonest wins; ties: first seen)
    spell = {}
    for tags in raw_tags.values():
        for t in tags:
            spell.setdefault(t.casefold(), Counter())[t] += 1
    canon = {k: max(c, key=lambda s: (c[s], -list(c).index(s))) for k, c in spell.items()}
    merged = {k: sorted(c) for k, c in spell.items() if len(c) > 1}

    # ---- join
    out, no_tags_row, untagged = [], [], []
    seen = set()
    for r in data:
        key = (r["word1"], r["word2"])
        seen.add(key)
        if key not in raw_tags:
            untagged.append(key)
            tags = []
        else:
            tags, used = [], set()
            for t in raw_tags[key]:
                c = canon[t.casefold()]
                if c not in used:                     # a tag repeated within one pair counts once
                    used.add(c); tags.append(c)
            if not tags:
                no_tags_row.append(key)
        row = {c: r.get(c, "") for c in OUT_COLUMNS if c in r or c in ("word1", "word2")}
        row[OUT_TAG_COLUMN] = TAG_JOIN.join(tags)
        out.append(row)
    not_in_data = [k for k in raw_tags if k not in seen]

    # ---- columns present in the data
    cols = [c for c in OUT_COLUMNS if any(c in r for r in out)] + [OUT_TAG_COLUMN]
    out_path = Path(args.out) if args.out else Path(args.data).with_name(Path(args.data).stem + "_tagged.csv")
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:      # BOM: Excel opens it correctly
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(out)

    # ---- report
    freq = Counter(t for r in out for t in r[OUT_TAG_COLUMN].split(TAG_JOIN) if t)
    n_with = sum(1 for r in out if r[OUT_TAG_COLUMN])
    print("Data file   : %d pairs (%s)" % (len(data), Path(args.data).name))
    print("Tagged file : %d rows (%s)" % (len(t_rows), Path(args.tagged).name))
    print("Pairs with tags: %d of %d;  %d different tags;  %.1f tags per pair on average"
          % (n_with, len(out), len(freq), sum(freq.values()) / max(n_with, 1)))
    if untagged:
        print("WARNING: %d pair(s) in the data file are not in the tagged file (left without tags): %s"
              % (len(untagged), ", ".join("%s+%s" % k for k in untagged[:10])))
    if not_in_data:
        print("WARNING: %d pair(s) in the tagged file are not in the data file and were not added: %s"
              % (len(not_in_data), ", ".join("%s+%s" % k for k in not_in_data[:10])))
    if no_tags_row:
        print("WARNING: empty %s for: %s" % (TAG_COL_TEXT, ", ".join("%s+%s" % k for k in no_tags_row[:10])))
    if dup:
        print("note: %d repeated pair(s) in the tagged file; the first one was used." % dup)
    for k, variants in sorted(merged.items()):
        print("note: spellings differing only in capitals were made the same: %s  ->  %s"
              %(" / ".join('"%s"' % v for v in variants), '"%s"' % canon[k]))
    print("\nTags (number of pairs):")
    for t, n in sorted(freq.items(), key=lambda x: (-x[1], x[0])):
        print("  %3d  %s" % (n, t))
    print("\nWrote %s" % out_path)


if __name__ == "__main__":
    main()
