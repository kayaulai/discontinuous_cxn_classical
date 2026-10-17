#!/usr/bin/env python3
"""
build_site.py -- turn an Excel (or CSV) table of word pairs into a finished,
fixed-data Collomap website (no "upload your own file" features).

    python build_site.py my_data.xlsx --config site.json --out my_site

Then open my_site/index.html in a browser, or upload the whole my_site folder to
any web host (GitHub Pages, a university web space, ...). It is plain HTML/CSS/JS:
no server and no database are needed.

Requirements: Python 3.8+ and, for .xlsx input,  pip install openpyxl

WHAT THE EXCEL FILE LOOKS LIKE
  Sheet "pairs"     one row per word pair (first sheet is used if there is no "pairs")
  Sheet "examples"  optional; one row per example sentence

  Pairs columns:    word1, word2        (required)
                    sym_score           number: mutual strength of the pair
                    asym_score          number: one-way pull of the pair (first word -> second word)
                    asym_score_2        optional number: the pull the other way (second word -> first);
                                        used when the second word is in focus. If missing, asym_score is used.
                    attr_w1_<name>      optional tag about the first word   (becomes a filter)
                    attr_w2_<name>      optional tag about the second word
                    attr_pair_<name>    optional tag about the pair
  Examples columns: word1, word2, text  (required)
                    translation, trans1, trans2, source_label, work, chapter, ctext_urn, ud_sent_id

  Your columns can be called anything: say so under "columns" and "tags" in the
  settings file (see example/site.json).

SETTINGS FILE (JSON; every key is optional)
  site_name         name shown in the header                     "Collomap"
  dataset_title     title of the dataset (also used in the citation)
  authors, year, url, license, source_url
                    citation / footer details
  about             a few sentences about the dataset ("\\n\\n" starts a new paragraph)
  group_names       {"first": "Adverb", "second": "Final particle"}
  default_distance  "asym" (one-way pull) or "sym" (mutual strength)
  default_spacing   "rank" (even steps), "raw" (true scale) or "log"
  sample_text       passage pre-filled in the "try it on your own text" box
  lang              language code of the page, e.g. "en" (default), "zh-Hant", "bo"
  font_family       font for the interface, e.g. "Helvetica, Arial, sans-serif" (default)
  word_font_family  font for the WORDS and sentences (list, graph labels, examples), e.g.
                    "'Noto Sans Tibetan', 'Microsoft Himalaya', sans-serif"; default = font_family
  word_line_height  line spacing for example sentences / own-text box (default 1.8; Tibetan
                    usually looks better at 2.2 or more)
  font_size         base text size in pixels, 12-24 (default 15)
  font_files        your own font files to ship with the site, so every visitor sees the same
                    letters: ["fonts/NotoSansTibetan.ttf"] or
                    [{"family": "Noto Sans Tibetan", "file": "fonts/NotoSansTibetan.ttf", "weight": "100 900"}]
                    (a loaded font is added to word_font_family automatically)
  colors            {"first": "#1a64a8", "second": "#b4500a", "accent": "#1a64a8"}
                    first/second = the two groups of words; accent = buttons and links
  sheet_pairs, sheet_examples
                    sheet names if yours differ from "pairs" / "examples"
  columns           {"Adverb": "word1", "Particle": "word2", "PPMI": "sym_score", ...}
                    maps YOUR column headings to the standard names (both sheets)
  tags              {"first": ["Semantic field"], "second": ["Position"], "pair": ["Register"]}
                    spreadsheet columns that should become filters
"""
import argparse, csv, io, json, re, shutil, sys, urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
CDN = {
    "cytoscape.min.js": "https://cdnjs.cloudflare.com/ajax/libs/cytoscape/3.28.1/cytoscape.min.js",
    "papaparse.min.js": "https://cdnjs.cloudflare.com/ajax/libs/PapaParse/5.4.1/papaparse.min.js",
}
PAIR_BASE = ["word1", "word2", "sym_score", "asym_score", "asym_score_2"]
EX_COLS = ["word1", "word2", "text", "translation", "trans1", "trans2", "source_label",
           "work", "chapter", "ctext_urn", "ud_sent_id", "w1_start", "w1_end", "w2_start", "w2_end"]

warnings = []
def warn(msg):
    warnings.append(msg); print("  WARNING: " + msg)
def die(msg):
    sys.exit("\nERROR: " + msg + "\n")

# ---------------------------------------------------------------- reading tables
def read_table(path, sheet=None, required_sheet=False):
    """Returns (rows as list of lists, sheet name used) or (None, None) if the sheet is absent."""
    path = Path(path)
    if not path.exists():
        die("Cannot find " + str(path))
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            die("Reading Excel files needs openpyxl.  Run:  pip install openpyxl")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        names = wb.sheetnames
        chosen = None
        if sheet:
            for n in names:
                if n.strip().lower() == sheet.strip().lower():
                    chosen = n
            if chosen is None and required_sheet:
                die('The workbook has no sheet called "%s". Sheets found: %s' % (sheet, ", ".join(names)))
        if chosen is None:
            if required_sheet is False and sheet is not None:
                return None, None
            chosen = names[0]
        rows = [list(r) for r in wb[chosen].iter_rows(values_only=True)]
        return rows, chosen
    if ext in (".csv", ".txt", ".tsv"):
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return list(csv.reader(io.StringIO(text), dialect)), path.name
    die("Unsupported file type %s (use .xlsx or .csv)" % ext)

def clean(v):
    if v is None: return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else format(v, ".10g")
    return str(v).strip()

def slug(s):
    return re.sub(r"[^0-9a-z\u00C0-\uFFFF]+", "_", s.strip().lower()).strip("_")

def to_dicts(rows, colmap, tags, label):
    """Header normalisation + the user's column renaming. Returns (headers, list of dicts)."""
    if not rows:
        die("The %s table is empty." % label)
    raw_headers = [clean(h) for h in rows[0]]
    cmap = {k.strip().lower(): v.strip() for k, v in (colmap or {}).items()}
    tagmap = {}
    for scope, prefix in (("first", "attr_w1_"), ("second", "attr_w2_"), ("pair", "attr_pair_")):
        for name in (tags or {}).get(scope, []):
            tagmap[name.strip().lower()] = prefix + slug(name)
    headers = []
    for h in raw_headers:
        key = h.lower()
        if key in cmap: headers.append(cmap[key].lower())
        elif key in tagmap: headers.append(tagmap[key])
        else: headers.append(key.replace(" ", "_") if key.startswith("attr_") else key)
    out = []
    for r in rows[1:]:
        cells = [clean(c) for c in r] + [""] * (len(headers) - len(r))
        if not any(cells): continue
        out.append({h: cells[i] for i, h in enumerate(headers) if h})
    return [h for h in headers if h], out, raw_headers

# ---------------------------------------------------------------- validating
def prepare_pairs(rows, cfg):
    headers, recs, raw = to_dicts(rows, cfg.get("columns"), cfg.get("tags"), "pairs")
    for need in ("word1", "word2"):
        if need not in headers:
            die('Could not find the "%s" column in the pairs sheet.\n  Columns found: %s\n'
                '  Either rename a column, or tell me which one it is in the settings file, e.g.\n'
                '  "columns": {"Adverb": "word1", "Particle": "word2", "PPMI": "sym_score"}'
                % (need, ", ".join(raw)))
    if not ({"sym_score", "asym_score", "asym_score_2"} & set(headers)):
        warn("No score columns (sym_score / asym_score) found: distances in the picture will not mean anything.")
    attr = [h for h in headers if h.startswith("attr_")]
    cols = [c for c in PAIR_BASE if c in headers] + attr
    kept, seen, nonnum = [], set(), 0
    skipped_blank = dupes = 0
    for r in recs:
        if not r.get("word1") or not r.get("word2"):
            skipped_blank += 1; continue
        key = (r["word1"], r["word2"])
        if key in seen:
            dupes += 1; continue
        seen.add(key)
        for sc in ("sym_score", "asym_score", "asym_score_2"):
            v = r.get(sc, "").replace(",", ".")      # tolerate decimal commas
            try:
                if v != "": float(v)
            except ValueError:
                nonnum += 1; v = ""
            if sc in r: r[sc] = v
        kept.append(r)
    if skipped_blank: warn("%d pair row(s) skipped because word1 or word2 was empty." % skipped_blank)
    if dupes: warn("%d duplicate word pair(s) ignored (the first occurrence of each is kept)." % dupes)
    if nonnum: warn("%d score value(s) were not numbers and were left blank." % nonnum)
    if not kept: die("No usable word pairs found in the pairs sheet.")
    return cols, kept

def prepare_examples(rows, cfg, pair_keys):
    headers, recs, raw = to_dicts(rows, cfg.get("columns"), None, "examples")
    for need in ("word1", "word2", "text"):
        if need not in headers:
            die('Could not find the "%s" column in the examples sheet.\n  Columns found: %s' % (need, ", ".join(raw)))
    cols = [c for c in EX_COLS if c in headers]
    kept, orphan = [], 0
    for r in recs:
        if not (r.get("word1") and r.get("word2") and r.get("text")): continue
        if (r["word1"].lower(), r["word2"].lower()) not in pair_keys:
            orphan += 1; continue
        kept.append(r)
    if orphan:
        warn("%d example sentence(s) skipped: their word pair is not in the pairs sheet." % orphan)
    return cols, kept

def to_csv(cols, recs):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(cols)
    for r in recs: w.writerow([r.get(c, "") for c in cols])
    return buf.getvalue()

# ---------------------------------------------------------------- appearance
def hex_to_rgb(h):
    h = h.lstrip("#")
    if len(h) == 3: h = "".join(c * 2 for c in h)
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))

def mix(rgb, other, t):
    return "#%02x%02x%02x" % tuple(round(c * (1 - t) + o * t) for c, o in zip(rgb, other))

def luminance(rgb):
    def f(c):
        c /= 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (f(c) for c in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b

def contrast_vs_white(rgb):
    return 1.05 / (luminance(rgb) + 0.05)

def with_generic(family):
    family = family.strip().rstrip(";")
    return family if re.search(r"(sans-serif|serif|monospace|cursive)\s*$", family, re.I) else family + ", sans-serif"

FONT_FORMATS = {".woff2": "woff2", ".woff": "woff", ".ttf": "truetype", ".otf": "opentype"}

def plan_appearance(cfg, base_dir):
    """Validates the look-and-feel settings. Returns (custom.css text, page language, [font files to copy])."""
    css = []
    root = []
    faces = []
    families = []
    copies = []
    for item in cfg.get("font_files") or []:
        spec = {"file": item} if isinstance(item, str) else dict(item)
        src = Path(spec.get("file", ""))
        if not src.is_absolute(): src = base_dir / src
        if not src.exists(): die("Font file not found: %s (paths in the settings file are relative to that file)" % src)
        fmt = FONT_FORMATS.get(src.suffix.lower())
        if not fmt: die("Unsupported font type %s (use .woff2, .woff, .ttf or .otf)" % src.suffix)
        fam = spec.get("family") or src.stem
        copies.append(src)
        faces.append("@font-face { font-family: %s; src: url('fonts/%s') format('%s'); font-weight: %s; font-style: %s; font-display: swap; }"
                     % (json.dumps(fam), src.name, fmt, spec.get("weight", "100 900"), spec.get("style", "normal")))
        if fam not in families: families.append(fam)

    ui = cfg.get("font_family")
    words = cfg.get("word_font_family")
    if ui: root.append("--font: %s;" % with_generic(ui))
    if words or families:
        base = with_generic(words or ui or "Helvetica, Arial, sans-serif")
        add = [json.dumps(f).replace('"', "'") for f in families if f.lower() not in base.lower()]
        if add:
            base = ", ".join(add) + ", " + base
            print("  fonts: %s added to the word font stack" % ", ".join(families))
        root.append("--word-font: %s;" % base)
    if cfg.get("word_line_height"):
        try: root.append("--word-lh: %s;" % float(cfg["word_line_height"]))
        except (TypeError, ValueError): die("word_line_height must be a number, e.g. 2.2")

    colors = cfg.get("colors") or {}
    for key, var in (("first", "c1"), ("second", "c2"), ("accent", "accent")):
        if key not in colors: continue
        val = str(colors[key]).strip().lower()
        if not re.fullmatch(r"#([0-9a-f]{3}|[0-9a-f]{6})", val):
            die('colors.%s must look like "#1a64a8" (got "%s")' % (key, colors[key]))
        rgb = hex_to_rgb(val)
        c = contrast_vs_white(rgb)
        if c < 4.5:
            warn("colors.%s (%s) has a contrast of only %.1f:1 against white; text in this colour may be hard to read (4.5:1 or more is recommended)." % (key, val, c))
        root.append("--%s: %s; --%s-dark: %s; --%s-dim: %s;" % (var, val, var, mix(rgb, (0, 0, 0), 0.25), var, mix(rgb, (255, 255, 255), 0.88)))
    unknown = set(colors) - {"first", "second", "accent"}
    if unknown: warn('Unknown colour name(s) ignored: %s (use "first", "second", "accent")' % ", ".join(sorted(unknown)))

    if cfg.get("font_size") not in (None, ""):
        try: fs = int(cfg["font_size"])
        except (TypeError, ValueError): die("font_size must be a whole number of pixels, e.g. 16")
        if not 12 <= fs <= 24: die("font_size must be between 12 and 24")
        css.append("body { font-size: %dpx; }" % fs)

    text = "/* Generated by build_site.py from your settings file. */\n" + "\n".join(faces)
    if root: text += "\n:root {\n  " + "\n  ".join(root) + "\n}"
    if css: text += "\n" + "\n".join(css)
    lang = str(cfg.get("lang") or "en").strip()
    if not re.fullmatch(r"[A-Za-z]{2,3}(-[A-Za-z0-9]{2,8})*", lang): die('lang must be a language code such as "en", "zh-Hant" or "bo"')
    return text + "\n", lang, copies

# ---------------------------------------------------------------- building
def js_json(obj):
    return json.dumps(obj, ensure_ascii=False, indent=1).replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")

def main():
    ap = argparse.ArgumentParser(description="Build a fixed-data Collomap website from an Excel/CSV file.",
                                 formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    ap.add_argument("data", help="Excel (.xlsx) or CSV file with the word pairs")
    ap.add_argument("--config", help="settings file (JSON)")
    ap.add_argument("--out", default="site", help="folder to create (default: site)")
    ap.add_argument("--examples", help="separate file with example sentences (if not a sheet in the workbook)")
    ap.add_argument("--template", default=str(HERE / "template"), help="folder holding index.html, style.css, app.js")
    ap.add_argument("--site-name"); ap.add_argument("--title", help="dataset title")
    ap.add_argument("--first-name", help='what to call the first group of words, e.g. "Adverb"')
    ap.add_argument("--second-name", help='what to call the second group, e.g. "Final particle"')
    ap.add_argument("--bundle-libs", action="store_true",
                    help="download Cytoscape and PapaParse into the site so it works without a CDN")
    ap.add_argument("--libs-dir", help="folder already holding cytoscape.min.js and papaparse.min.js to bundle")
    a = ap.parse_args()

    cfg = {}
    if a.config:
        try:
            cfg = json.loads(Path(a.config).read_text(encoding="utf-8-sig"))
        except (OSError, ValueError) as e:
            die("Could not read the settings file %s: %s" % (a.config, e))
    if a.site_name: cfg["site_name"] = a.site_name
    if a.title: cfg["dataset_title"] = a.title
    if a.first_name or a.second_name:
        g = dict(cfg.get("group_names") or {})
        if a.first_name: g["first"] = a.first_name
        if a.second_name: g["second"] = a.second_name
        cfg["group_names"] = g
    known = {"site_name", "dataset_title", "authors", "year", "url", "license", "source_url", "about", "group_names",
             "default_distance", "default_spacing", "sample_text", "sheet_pairs", "sheet_examples", "columns", "tags",
             "lang", "font_family", "word_font_family", "word_line_height", "font_size", "font_files", "colors"}
    for k in cfg:
        if k not in known and not k.startswith("_"):
            warn('Unknown setting "%s" ignored (typo?). Known settings: %s' % (k, ", ".join(sorted(known))))
    if cfg.get("default_distance") not in (None, "", "sym", "asym"): die('default_distance must be "sym" or "asym"')
    if cfg.get("default_spacing") not in (None, "", "rank", "raw", "log"): die('default_spacing must be "rank", "raw" or "log"')

    base_dir = Path(a.config).resolve().parent if a.config else Path.cwd()
    custom_css, lang, font_copies = plan_appearance(cfg, base_dir)   # validate the look-and-feel settings first

    tpl = Path(a.template)
    for f in ("index.html", "style.css", "app.js"):
        if not (tpl / f).exists(): die("Template file missing: %s" % (tpl / f))

    print("Reading", a.data)
    prow, psheet = read_table(a.data, cfg.get("sheet_pairs", "pairs"), required_sheet=False)
    if prow is None:
        prow, psheet = read_table(a.data, None)
    print("  pairs from sheet/file: %s" % psheet)
    pcols, pairs = prepare_pairs(prow, cfg)

    ex_rows = None
    if a.examples:
        ex_rows, esheet = read_table(a.examples, cfg.get("sheet_examples", "examples"), required_sheet=False)
        if ex_rows is None: ex_rows, esheet = read_table(a.examples, None)
    elif Path(a.data).suffix.lower() in (".xlsx", ".xlsm"):
        ex_rows, esheet = read_table(a.data, cfg.get("sheet_examples", "examples"), required_sheet=False)
        if esheet is not None and esheet == psheet: ex_rows = None
    pair_keys = {(r["word1"].lower(), r["word2"].lower()) for r in pairs}
    ecols, examples = ([], [])
    if ex_rows:
        ecols, examples = prepare_examples(ex_rows, cfg, pair_keys)

    # ---- summary ----
    w1 = {r["word1"] for r in pairs}; w2 = {r["word2"] for r in pairs}
    tagcols = [c for c in pcols if c.startswith("attr_")]
    print("\nData summary")
    print("  %d word pairs; %d different first words, %d different second words" % (len(pairs), len(w1), len(w2)))
    print("  tag columns (filters): %s" % (", ".join(tagcols) if tagcols else "none"))
    print("  example sentences: %d (for %d pairs)" % (len(examples), len({(e["word1"].lower(), e["word2"].lower()) for e in examples})))
    if len(w1) * len(w2) > 4e6:
        warn("This is a large dataset (%d x %d words); the overview layout may be slow in the browser." % (len(w1), len(w2)))

    # ---- write the site ----
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    html = (tpl / "index.html").read_text(encoding="utf-8")
    name = cfg.get("site_name") or "Collomap"
    html = html.replace('<body class="no-data">', '<body class="no-data fixed">')
    html = re.sub(r"<title>.*?</title>", "<title>%s – %s</title>" % (
        name.replace("<", "&lt;"), (cfg.get("dataset_title") or "explore which words go together").replace("<", "&lt;")), html, count=1)
    html = html.replace('<script src="app.js"></script>', '<script src="data.js"></script>\n<script src="app.js"></script>')

    if a.bundle_libs or a.libs_dir:
        (out / "lib").mkdir(exist_ok=True)
        ok = True
        for fn, url in CDN.items():
            dest = out / "lib" / fn
            try:
                if a.libs_dir: shutil.copyfile(Path(a.libs_dir) / fn, dest)
                else:
                    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
                    dest.write_bytes(urllib.request.urlopen(req, timeout=30).read())
            except Exception as e:
                ok = False; warn("Could not bundle %s (%s); the site will load it from the CDN instead." % (fn, e))
        if ok:
            for fn, url in CDN.items():
                html = html.replace(url, "lib/" + fn)
    shutil.copyfile(tpl / "style.css", out / "style.css")
    (out / "custom.css").write_text(custom_css, encoding="utf-8")
    if font_copies: (out / "fonts").mkdir(exist_ok=True)
    for src in font_copies: shutil.copyfile(src, out / "fonts" / src.name)
    html = html.replace('<link rel="stylesheet" href="style.css">', '<link rel="stylesheet" href="style.css">\n<link rel="stylesheet" href="custom.css">')
    html = html.replace('<html lang="en">', '<html lang="%s">' % lang, 1)
    (out / "index.html").write_text(html, encoding="utf-8")
    shutil.copyfile(tpl / "app.js", out / "app.js")

    site = {"siteName": name, "datasetTitle": cfg.get("dataset_title", ""), "authors": cfg.get("authors", ""),
            "year": str(cfg.get("year", "") or ""), "url": cfg.get("url", ""), "license": cfg.get("license", ""),
            "sourceUrl": cfg.get("source_url", ""), "about": cfg.get("about", ""),
            "groupNames": cfg.get("group_names") or None, "defaultDistance": cfg.get("default_distance", ""),
            "defaultSpacing": cfg.get("default_spacing", ""), "sampleText": cfg.get("sample_text", "")}
    data = {"pairsCsv": to_csv(pcols, pairs), "examplesCsv": to_csv(ecols, examples) if examples else ""}
    (out / "data.js").write_text(
        "/* Generated by build_site.py -- edit the Excel file / settings and rebuild instead of editing this. */\n"
        "window.COLLOMAP_SITE = %s;\nwindow.COLLOMAP_DATA = %s;\n" % (js_json(site), js_json(data)), encoding="utf-8")

    print("\nDone. Site written to: %s" % out.resolve())
    print("  Open %s in a browser to look at it," % (out / "index.html"))
    print("  or upload the whole folder to a web host.")
    if warnings: print("  (%d warning(s) above -- worth a look.)" % len(warnings))

if __name__ == "__main__":
    main()
