#!/usr/bin/env python3
"""
convert_exp2.py  --  turn the experiment-2 (pre-collation) results into the Excel file the site builder reads.

    python convert_exp2.py exp2_selected_df.csv

reads  the selected file from src/03_precollation.Rmd (one row per selected pair, with the measures
       in the same file: dc, sfp, n, pmi, kld_norm_dc_cond_sfp, kld_norm_sfp_cond_dc).
       Unlike experiment 1 there is no separate measures file.  Adverbs and particles may be
       collated units of two characters (e.g. 豈不, 而已).
writes exp2_site_data.xlsx   sheet "pairs": Non-final particle, Final particle, PMI, the two one-way pulls, Frequency
       exp2_site.json        the settings file for build_site.py (names, columns, fonts, colours)

and prints the command that builds the website:

    python build_site.py exp2_site_data.xlsx --config exp2_site.json --out exp2_site

The headings are the same as convert_exp1.py's, so add_tags.py works on the result:

    python add_tags.py exp2_site_data.xlsx exp2_selected_df_tagged.csv

Change anything in the block below (column names, labels, output names) and run it again.
Needs:  pip install openpyxl     (nothing else; no pandas)
"""
import argparse, json
from pathlib import Path

from convert_exp1 import die, read_csv, need, number

# ============================================================ SETTINGS: edit here
# the two slots, as they are called on the site
SLOT1_NAME = "Non-final particle"            # first word of a pair  (column "dc")
SLOT2_NAME = "Final particle"             # second word of a pair (column "sfp")

# columns of the selected file
COL_SLOT1 = "dc"
COL_SLOT2 = "sfp"
COL_SYM = "pmi"                       # mutual strength (symmetric)
# One-way pulls, read the same way as in convert_exp1.py:
#   kld_norm_sfp_cond_dc  = how much knowing the adverb narrows down the particle   (NFP -> FP)
#   kld_norm_dc_cond_sfp  = how much knowing the particle narrows down the adverb   (FP -> NFP)
COL_ASYM_1TO2 = "kld_norm_sfp_cond_dc"
COL_ASYM_2TO1 = "kld_norm_dc_cond_sfp"
COL_FREQ = "n"                        # raw co-occurrence count, kept in the Excel file for reference

# headings used in the Excel file
HEAD_SYM = "PMI"
HEAD_ASYM_1TO2 = "%s -> %s" % (SLOT1_NAME, SLOT2_NAME)
HEAD_ASYM_2TO1 = "%s -> %s" % (SLOT2_NAME, SLOT1_NAME)
HEAD_FREQ = "Frequency"

# the website's own settings (written to the .json file; see README.md for all options)
SITE_SETTINGS = {
    "site_name": "Collomap",
    "dataset_title": "Mapping the particle frames of Ancient Chinese: Statistical extraction and semantic analysis of an elusive discontinuous construction - Experiment 2",
    "authors": "Lai, Ryan Ka Yau",
    "year": "2026",
    "url": "",
    "license": "",
    "source_url": "https://github.com/kayaulai/discontinuous_cxn_classical",
    "about": ("Non-final and final particles in Late Archaic Chinese, after collocations among "
              "non-final particles and among final particles (such as 豈不 or 而已) were first "
              "merged into single units. "
              "Mutual strength is pointwise mutual information (PMI). The one-way pulls are "
              "normalised Kullback-Leibler divergences: when you look at a non-final particle, the "
              "distance shows how strongly it narrows down which final particle follows; when you "
              "look at a final particle, how strongly it narrows down the non-final particle."),
    "group_names": {"first": SLOT1_NAME, "second": SLOT2_NAME},
    "default_distance": "asym",
    "default_spacing": "rank",
    "sample_text": "有朋自遠方來，不亦樂乎？",
    "lang": "zh-Hant",
    "font_family": "Helvetica, Arial, sans-serif",
    "word_font_family": "'Noto Sans CJK TC', 'PingFang TC', 'Microsoft JhengHei', 'Source Han Sans TC', sans-serif",
    "word_line_height": 1.8,
    "colors": {"first": "#1a64a8", "second": "#b4500a", "accent": "#1a64a8"},
}
# ============================================================ end of settings


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("selected", help="exp2_selected_df.csv")
    ap.add_argument("--out-dir", default=".", help="folder for the two output files (default: here)")
    ap.add_argument("--name", default="exp2_site", help="start of the output file names (default: exp2_site)")
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        die("The openpyxl package is needed to write the Excel file:   pip install openpyxl")

    s_cols, s_rows = read_csv(args.selected, "Selected file")
    need(s_cols, [COL_SLOT1, COL_SLOT2, COL_SYM, COL_ASYM_1TO2, COL_ASYM_2TO1], "The selected file", args.selected)

    # --- walk the selection, in its own order
    out, seen = [], set()
    dup_sel, blank_sel = 0, 0
    gaps = {HEAD_SYM: 0, HEAD_ASYM_1TO2: 0, HEAD_ASYM_2TO1: 0}
    for r in s_rows:
        a, f = r[COL_SLOT1].strip(), r[COL_SLOT2].strip()
        if not a or not f:
            blank_sel += 1
            continue
        if (a, f) in seen:
            dup_sel += 1
            continue
        seen.add((a, f))
        vals = [number(r[COL_SYM]), number(r[COL_ASYM_1TO2]), number(r[COL_ASYM_2TO1])]
        for v, h in zip(vals, gaps):
            if v is None: gaps[h] += 1
        freq = number(r.get(COL_FREQ, ""))
        out.append([a, f] + vals + [int(freq) if freq is not None and freq == int(freq) else freq])

    if not out:
        die("No usable pairs in %s." % args.selected)

    # --- report
    print("Selected file : %d rows (%s)" % (len(s_rows), Path(args.selected).name))
    print("Pairs written : %d   (%d %s, %d %s)" % (
        len(out), len({o[0] for o in out}), SLOT1_NAME, len({o[1] for o in out}), SLOT2_NAME))
    multi = sorted({w for o in out for w in o[:2] if len(w) > 1})
    if multi:
        print("Collated units: %s" % " ".join(multi))
    if dup_sel:   print("note: %d repeated row(s) in the selected file were counted once." % dup_sel)
    if blank_sel: print("note: %d row(s) in the selected file with an empty adverb or particle were skipped." % blank_sel)
    for h, n in gaps.items():
        if n: print("WARNING: %d pair(s) have no usable value for %s (left blank)." % (n, h))

    # --- Excel file
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / (args.name + "_data.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "pairs"
    ws.append([SLOT1_NAME, SLOT2_NAME, HEAD_SYM, HEAD_ASYM_1TO2, HEAD_ASYM_2TO1, HEAD_FREQ])
    for row in out:
        ws.append(row)
    for c in ws[1]:
        c.font = openpyxl.styles.Font(bold=True)
    ws.freeze_panes = "A2"
    for col, w in zip("ABCDEF", (10, 10, 10, 12, 12, 11)):
        ws.column_dimensions[col].width = w
    for r in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for c in r: c.number_format = "0.000"
    wb.save(xlsx)

    # --- settings file for the builder
    cfg = dict(SITE_SETTINGS)
    cfg["columns"] = {SLOT1_NAME: "word1", SLOT2_NAME: "word2", HEAD_SYM: "sym_score",
                      HEAD_ASYM_1TO2: "asym_score", HEAD_ASYM_2TO1: "asym_score_2"}
    js = out_dir / (args.name + ".json")
    with open(js, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)
        f.write("\n")

    print("\nWrote %s\n      %s" % (xlsx, js))
    print("\nNext, build the website:\n  python build_site.py %s --config %s --out %s" % (
        xlsx, js, out_dir / args.name))


if __name__ == "__main__":
    main()
